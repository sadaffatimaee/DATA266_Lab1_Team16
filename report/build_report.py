import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
P1 = ROOT / "task1_llm" / "Poushali_Purkayastha"
P2 = ROOT / "task2_sentiment" / "Poushali_Purkayastha"
P3 = ROOT / "task3_gan" / "Poushali_Purkayastha"
S1 = ROOT / "task1_llm" / "Sadaf_Fatima_Syeda"
S2 = ROOT / "task2_sentiment" / "Sadaf_Fatima_Syeda"
S3 = ROOT / "task3_gan" / "Sadaf_Fatima_Syeda"
OUT = ROOT / "report" / "DATA266_Lab1_Report_Team_16.docx"
REPO_URL = "https://github.com/sadaffatimaee/DATA266_Lab1_Team16"


def kv_csv(path):
    with open(path, encoding="utf-8", newline="") as f:
        rows = list(csv.reader(f))
    return {r[0]: r[1] for r in rows[1:] if len(r) >= 2}


def wide_csv(path):
    with open(path, encoding="utf-8", newline="") as f:
        rows = list(csv.reader(f))
    header = rows[0]
    return {r[0]: dict(zip(header[1:], r[1:])) for r in rows[1:]}


def model_rows_csv(path):
    with open(path, encoding="utf-8", newline="") as f:
        rows = list(csv.DictReader(f))
    return {r["Model"]: r for r in rows}


def group_csv(path):
    with open(path, encoding="utf-8", newline="") as f:
        rows = list(csv.DictReader(f))
    return {r["metric"]: r["value"] for r in rows}


def fmt(v, nd=4):
    try:
        x = float(v)
    except (TypeError, ValueError):
        return str(v)
    if x != x:
        return "nan"
    if abs(x) >= 1000:
        return f"{x:,.0f}"
    if abs(x) >= 100:
        return f"{x:.1f}"
    if abs(x) < 1e-3 and x != 0:
        return f"{x:.2e}"
    return f"{x:.{nd}f}"


def pct(v):
    return f"{100 * float(v):.1f}%"


class Report:
    def __init__(self):
        self.doc = Document()
        st = self.doc.styles["Normal"]
        st.font.name = "Calibri"
        st.font.size = Pt(10)
        for s in self.doc.sections:
            s.left_margin = s.right_margin = Inches(0.8)
            s.top_margin = s.bottom_margin = Inches(0.8)

    def h(self, text, level=1):
        self.doc.add_heading(text, level=level)

    def p(self, text):
        self.doc.add_paragraph(text)

    def bullets(self, items):
        for it in items:
            self.doc.add_paragraph(it, style="List Bullet")

    def table(self, header, rows, font=8, widths=None):
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        for i, h in enumerate(header):
            t.rows[0].cells[i].text = str(h)
        for r in rows:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = "" if v is None else str(v)
        for row in t.rows:
            for i, c in enumerate(row.cells):
                if widths and i < len(widths):
                    c.width = Inches(widths[i])
                for par in c.paragraphs:
                    par.paragraph_format.space_after = Pt(0)
                    for run in par.runs:
                        run.font.size = Pt(font)
        for c in t.rows[0].cells:
            for par in c.paragraphs:
                for run in par.runs:
                    run.font.bold = True
        self.doc.add_paragraph()
        return t

    def picture(self, path, width=6.0, caption=None):
        path = Path(path)
        if path.exists():
            self.doc.add_picture(str(path), width=Inches(width))
            if caption:
                c = self.doc.add_paragraph(caption)
                for run in c.runs:
                    run.font.size = Pt(8)
                    run.font.italic = True
        else:
            self.p(f"[missing image: {path.relative_to(ROOT)}]")

    def landscape(self):
        sec = self.doc.add_section()
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
        return sec

    def portrait(self):
        sec = self.doc.add_section()
        sec.orientation = WD_ORIENT.PORTRAIT
        if sec.page_width > sec.page_height:
            sec.page_width, sec.page_height = sec.page_height, sec.page_width
        return sec


def build():
    r = Report()
    d = r.doc
    d.add_heading("DATA 266 Lab 1 Report, Team 16", 0)
    r.p("LLM pretraining from scratch, Yelp polarity sentiment classification, CycleGAN style transfer")
    r.p("Poushali Purkayastha and Sadaf Fatima Syeda. Fall 2026. Report date 2026-09-28.")
    r.p(f"Repository: {REPO_URL}")
    r.p("Kaggle team: PairProgramming_Team_16, competition data-266-fall-2026-gan-image-style-transfer.")

    r.h("1. Team ownership statement", 1)
    r.p(
        "Each member independently designed, coded, trained, and documented their own model for all three tasks, under their own named folder in the shared repository. "
        "Poushali Purkayastha built a 4-layer character-level GPT (Task 1); a mean-pooled embedding MLP baseline, a TextCNN, and a BiLSTM (Task 2); and a ResNet-6 CycleGAN with least-squares loss and identity loss (Task 3), together with the repository layout, the README with the one-command smoke test, the run logging and manifest utility, the Task 1 and Task 2 evaluation scripts, the competition FID and MiFID evaluator, and the blinded audit tooling in her folders. "
        "Sadaf Fatima Syeda built a 6-layer character-level GPT (Task 1); a max-pooled embedding baseline, a BiGRU with attention pooling, and a from-scratch Transformer encoder (Task 2); and a U-Net CycleGAN with sigmoid cross-entropy loss (Task 3), together with the shared TinyStories data pointer, her own evaluation and audit tooling, and the Kaggle team. "
        "The comparison tables and the joint analyses in this report were written together; each member's results.md and failure_analysis.md are their own."
    )

    r.h("2. Repository, reproducibility, and hardware", 1)
    r.bullets([
        "Layout follows the lab brief: task1_llm, task2_sentiment, task3_gan, each with a data pointer and one folder per member holding src (code and the executed notebook), data_processed, checkpoints, outputs, metrics_report.csv, failure_analysis.md, and results.md; reproducibility/raw_logs and reproducibility/manifests; this report in report/.",
        "One-command smoke test from a clean clone, documented in the README: python task1_llm/Poushali_Purkayastha/src/run.py --config configs/smoke.yaml. Every run is driven by a YAML or JSON config; no personal paths or secrets are committed.",
        "Every training run wrote an untouched raw log and a manifest with the config, hardware, package versions, and the checkpoint that produced each reported number. Poushali's final runs: task1_llm_full_20260928_180129, task2_sentiment_full_20260928_180234, task3_gan_full_20260928_180436, all on an NVIDIA GeForce RTX 5090 (CUDA 12.8, torch 2.11.0+cu128, bfloat16) in the GPU lab. Sadaf's final runs: task1_sadaf_20260926_050656, task2_sadaf_20260926_061012, task3_sadaf_20260927_181844, all on a Tesla T4 in Google Colab.",
        "Raw datasets are not committed. Zipped copies with read access are linked in the README's Datasets table (TinyStories archive, Yelp polarity parquet files, the competition's monet_jpg and photo_jpg).",
        "Generated Task 3 images, the Kaggle image archive, and per-epoch checkpoints are backed up on Google Drive; the committed checkpoints are the ones that produced every reported number.",
    ])

    # ---------------- Task 1 ----------------
    p1 = kv_csv(P1 / "metrics_report.csv")
    s1 = kv_csv(S1 / "metrics_report.csv")
    r.h("3. Task 1: GPT-style LLM from scratch on TinyStories", 1)
    r.h("3.1 Comparison table", 2)
    rows = [
        ["Architecture", "4 pre-norm Transformer blocks, 4 heads of 64 dims, d_model 256, FFN 1024 GELU, own LayerNorm, causal self-attention from scratch, learnable token and positional embeddings, linear LM head", "6 Transformer blocks, 6 heads, d_model 192, ReLU, dropout 0.2, causal self-attention from scratch, learnable embeddings, linear LM head"],
        ["Context length (characters)", "128", "256"],
        ["Vocabulary", "93 characters plus unknown", "91 characters"],
        ["Training data", "100,000 windows of 128 characters (12.8M tokens) from 15,108 stories; 10,000 validation windows from 1,678 held-out stories", "100K characters of TinyStories for training, 10K characters for validation"],
        ["Optimizer and schedule", "AdamW lr 6e-4, betas 0.9/0.95, weight decay 0.1, 300-step warm-up, cosine decay to 6e-5, clip 1.0", "AdamW lr 6e-4, warm-up then linear decay, clip 1.0"],
        ["Batch, epochs, steps", "128 sequences, 10 epochs, 7,820 steps", "32 sequences, 10 epochs, 7,800 steps"],
        ["Mixed precision, hardware", "bfloat16, NVIDIA GeForce RTX 5090", "Tesla T4 (Colab)"],
        ["Parameter count", fmt(p1["Parameter Count"]), fmt(s1["Parameter Count"])],
        ["Training cross-entropy (final)", fmt(p1["Training Cross-Entropy Loss"]), fmt(s1["Training Cross-Entropy Loss"])],
        ["Validation cross-entropy (final)", fmt(p1["Validation Cross-Entropy Loss"]), fmt(s1["Validation Cross-Entropy Loss"])],
        ["Best validation cross-entropy", fmt(p1["Best Validation Cross-Entropy Loss"]) + " (epoch 10)", fmt(s1["Best Validation Loss"]) + " (epoch 2)"],
        ["Perplexity (per character)", fmt(p1["Perplexity"]), fmt(s1["Perplexity"])],
        ["Bits-per-character", fmt(p1["Bits-Per-Character"]), fmt(s1["Bits-Per-Character (BPC)"])],
        ["Generalization gap (val - train)", fmt(p1["Generalization Gap"]), fmt(s1["Generalization Gap (Val - Train)"])],
        ["Top-1 next-character accuracy", pct(p1["Top-1 Next-Character Accuracy"]), pct(s1["Top-1 Next-Character Accuracy"])],
        ["Distinct-1 / 2 / 3", f'{fmt(p1["Distinct-1"], 3)} / {fmt(p1["Distinct-2"], 3)} / {fmt(p1["Distinct-3"], 3)}', f'{fmt(s1["Distinct-1 (temp sampling)"], 3)} / {fmt(s1["Distinct-2 (temp sampling)"], 3)} / {fmt(s1["Distinct-3 (temp sampling)"], 3)}'],
        ["Repeated 4-gram rate", fmt(p1["Repeated 4-gram Rate"], 3), f'{fmt(s1["Repeated 4-gram Rate (temp sampling)"], 3)} (sampling), {fmt(s1["Repeated 4-gram Rate (greedy)"], 3)} (greedy)'],
        ["Gradient norm mean / max", f'{fmt(p1["Gradient Norm Mean"], 3)} / {fmt(p1["Gradient Norm Max"], 2)}', f'{fmt(s1["Gradient Norm Mean"], 3)} / {fmt(s1["Gradient Norm Max"], 2)}'],
        ["Loss spikes / NaN losses", f'{p1["Loss Spikes"]} / {p1["NaN or Inf Losses"]}', f'{s1["Loss Spikes"]} / {s1["NaN/Inf Losses"]}'],
        ["Training tokens/sec", fmt(p1["Training Tokens/sec"]), fmt(s1["Training Throughput (tokens/sec)"])],
        ["Generation tokens/sec", fmt(p1["Generation Tokens/sec"], 1), fmt(s1["Generation Speed (tokens/sec)"], 1)],
        ["Peak memory (MB)", fmt(p1["Peak Memory (MB)"], 1), fmt(s1["Peak Memory Usage (MB)"], 1)],
        ["Total training time (sec)", fmt(p1["Total Training Time (sec)"], 1), fmt(s1["Total Training Time (sec)"], 1)],
        ["Evidence", "task1_llm/Poushali_Purkayastha/outputs/full/loss_curves.png, training_dynamics.png, samples.txt, history.json; checkpoints/full/final.pt; raw log and manifest Poushali_Purkayastha_task1_llm_full_20260928_180129", "task1_llm/Sadaf_Fatima_Syeda/outputs/loss_curves.png, training_diagnostics.png, generated_samples.txt; checkpoints/best_model.pt; raw log and manifest task1_sadaf_20260926_050656"],
    ]
    r.table(["", "Poushali Purkayastha", "Sadaf Fatima Syeda"], rows, font=8, widths=[1.6, 2.6, 2.6])
    r.picture(P1 / "outputs" / "full" / "loss_curves.png", 6.0, "Poushali: training and validation loss per epoch and per step, 10 epochs.")
    r.picture(S1 / "outputs" / "loss_curves.png", 6.0, "Sadaf: training and validation loss curves.")

    r.h("3.2 Joint analysis", 2)
    r.p(
        "Strengths. Both models are complete from-scratch decoders with causal masking, learnable positional embeddings, warm-up and a decaying schedule, and both trained without a single loss spike or NaN. "
        "Poushali's model reaches a validation loss of 0.748 nats per character (1.08 bits per character, 76.3% top-1 accuracy) and its validation loss was still falling at epoch 10, so it is under-trained rather than over-fitted. "
        "Sadaf's deeper 6-layer model reached its best validation loss of 1.314 at epoch 2 and generates fluent TinyStories openings with greedy decoding."
    )
    r.p(
        "Weaknesses. The two runs used different amounts of text: Poushali's split is 100,000 windows of 128 characters (12.8 million characters), Sadaf's is a 100,000-character training text. With 128 times less data the 6-layer model memorizes the training text after two epochs (final training loss 0.055 against validation 2.609, a gap of 2.55) and its samples repeat training passages, which is the memorization failure in her analysis. Poushali's larger split avoids overfitting but her greedy samples still fall into repetition loops and the 128-character context cannot keep characters consistent across a paragraph."
    )
    r.p(
        "Limitations. Character-level modelling caps the coherence any of these models can reach: neither model has a notion of words, so high-temperature samples contain invented words, and neither context window covers a whole story. Diversity metrics are computed on a handful of samples and are noisy. Throughput numbers are not comparable across the two members because the hardware differs (RTX 5090 against T4)."
    )
    r.p(
        "What the team would try next. Train both models on the same 100,000-window split so depth and context are the only differences; add early stopping on validation loss; raise the context to 256 or 512 characters for the 4-layer model; use nucleus sampling with a repetition penalty at inference; and, beyond the brief, compare a small word-piece tokenizer against characters at equal parameter count."
    )

    r.h("3.3 Failure analyses", 2)
    r.p("Poushali Purkayastha, from task1_llm/Poushali_Purkayastha/failure_analysis.md; snippets are verbatim from outputs/full/samples.txt.")
    r.table(["Case", "Snippet", "Failure type", "Observation"], [
        ["1, greedy, prompt 'One day, a little girl'", "She said, \"Lily, you can have a ball too. It is my ball. It is my ball. It is my ball. It is my ball. It is my ball. It is my ball. It has a ball. It can make a ball.", "Repetition", "Greedy decoding locks onto the highest-probability continuation of its own output; no temperature sample loops this way and the repeated 4-gram rate over all samples is 5%."],
        ["2, temperature 1.0, prompt 'The cat'", "He wanted to unlect his ordinary too. [...] so his mom roared and his car cra. We could say fast!", "Broken grammar, hallucinated tokens", "A character model can sample a plausible-looking letter string that is not a word, or end a word early; only the temperature 1.0 samples show it."],
        ["3, temperature 0.7, prompt 'One day, a little girl'", "One day, a little girl named Lily went outside and saw a little girl named Lily. [...] They had many ducks and tasted on the road. [...] Mom said they lost their mom and dad.", "Loss of coherence", "Every sentence is grammatical but the story contradicts itself; the 128-character context holds about two sentences, so entities are not tracked across a paragraph."],
    ], font=8, widths=[1.3, 2.6, 1.0, 2.0])
    r.p("Sadaf Fatima Syeda, from task1_llm/Sadaf_Fatima_Syeda/failure_analysis.md; snippets from outputs/generated_samples.txt.")
    r.table(["Case", "Snippet", "Failure type", "Observation"], [
        ["1, greedy", "Once upon a time, there was a little girl named Lily. She loved to eat blueberries. One day, Lily went to the park with her mom. There, she found a big blueberry bush with lots of sweet blueberries.", "Memorisation / overfitting", "The nearly perfect continuation suggests the model repeats text from the training data rather than generating new text."],
        ["2, temperature 0.8, seed 2", "\"What are you doing on my bake it!\" Tim mom said, \"I have a special poor my for the bird.\"", "Broken grammar", "The words are recognizable, but the character-level model produces broken sentence structure and grammatically incorrect phrases."],
        ["3, temperature 0.8, seeds 0 and 2", "Inside the hut, they saw a big tree. On the tree, there was a swing. / So, I wanted to fold the paper to make a boat.\" Her mom smiled and said, \"That's better. I love you both.\"", "Loss of coherence / hallucination", "The story places a tree inside a hut and becomes unclear about the characters' identities and actions."],
    ], font=8, widths=[1.3, 2.6, 1.0, 2.0])

    # ---------------- Task 2 ----------------
    pw = wide_csv(P2 / "outputs" / "full" / "metrics_wide.csv")
    sm = model_rows_csv(S2 / "metrics_report.csv")
    sb, sg, stf = sm["Baseline"], sm["BiGRU_Attention"], sm["Transformer_Encoder"]
    r.landscape()
    r.h("4. Task 2: Yelp polarity sentiment classification", 1)
    r.h("4.1 Comparison table, all six models", 2)

    def pm(metric, model, nd=4):
        return fmt(pw[metric][model], nd)

    def sv(row, key, nd=4):
        v = row.get(key, "")
        return fmt(v, nd) if v not in ("", None) else ""

    header = ["", "P: baseline", "P: textcnn", "P: bilstm", "S: Baseline", "S: BiGRU + attention", "S: Transformer encoder"]
    rows = [
        ["Architecture", "learned 128-d embedding, masked mean pooling, MLP 128, dropout 0.3", "learned 128-d embedding, Conv1d widths 3/4/5 x 100 filters, masked max-over-time, dropout 0.5", "learned 128-d embedding, BiLSTM 128 per direction over packed sequences, final states, dropout 0.3", "learned 128-d embedding, max pooling, linear layer", "learned 128-d embedding, BiGRU hidden 64, attention pooling", "learned 128-d embedding, 4 heads, 2 layers, dropout 0.1, own attention code from Task 1"],
        ["Preprocessing", "lowercase, punctuation removed, NLTK stopwords minus negations, Snowball stemming, 20K vocab, 256 tokens", "same", "same", "lowercase, links and punctuation removed, stopwords minus negations, lemmatization, 20K vocab, 256 tokens", "same", "same"],
        ["Training data", "100K reviews (seed 20), 10K validation, official 38K test", "same", "same", "100K reviews (seed 266), 5K validation, official 38K test", "same", "same"],
        ["Optimizer, batch, epochs", "Adam 1e-3, batch 256, 4 epochs, best by val macro-F1", "same", "same", "Adam 1e-3, batch 64, 3 epochs, best by val accuracy", "same", "Adam 5e-4, batch 64, 3 epochs"],
        ["Hardware", "RTX 5090, bfloat16", "RTX 5090, bfloat16", "RTX 5090, bfloat16", "Tesla T4", "Tesla T4", "Tesla T4"],
        ["Parameter count", pm("Parameter Count", "baseline"), pm("Parameter Count", "textcnn"), pm("Parameter Count", "bilstm"), sv(sb, "Params"), sv(sg, "Params"), sv(stf, "Params")],
        ["Accuracy", pm("Accuracy", "baseline"), pm("Accuracy", "textcnn"), pm("Accuracy", "bilstm"), sv(sb, "Accuracy"), sv(sg, "Accuracy"), sv(stf, "Accuracy")],
        ["Accuracy 95% CI", f'{pm("Accuracy 95% CI low", "baseline")} to {pm("Accuracy 95% CI high", "baseline")}', f'{pm("Accuracy 95% CI low", "textcnn")} to {pm("Accuracy 95% CI high", "textcnn")}', f'{pm("Accuracy 95% CI low", "bilstm")} to {pm("Accuracy 95% CI high", "bilstm")}', f'{sv(sb, "Accuracy_CI95_low")} to {sv(sb, "Accuracy_CI95_high")}', f'{sv(sg, "Accuracy_CI95_low")} to {sv(sg, "Accuracy_CI95_high")}', f'{sv(stf, "Accuracy_CI95_low")} to {sv(stf, "Accuracy_CI95_high")}'],
        ["Precision / recall / F1, macro", f'{pm("Precision (macro)", "baseline")} / {pm("Recall (macro)", "baseline")} / {pm("F1 (macro)", "baseline")}', f'{pm("Precision (macro)", "textcnn")} / {pm("Recall (macro)", "textcnn")} / {pm("F1 (macro)", "textcnn")}', f'{pm("Precision (macro)", "bilstm")} / {pm("Recall (macro)", "bilstm")} / {pm("F1 (macro)", "bilstm")}', f'{sv(sb, "Precision_macro")} / {sv(sb, "Recall_macro")} / {sv(sb, "F1_macro")}', f'{sv(sg, "Precision_macro")} / {sv(sg, "Recall_macro")} / {sv(sg, "F1_macro")}', f'{sv(stf, "Precision_macro")} / {sv(stf, "Recall_macro")} / {sv(stf, "F1_macro")}'],
        ["Precision / recall / F1, micro", f'{pm("Precision (micro)", "baseline")} / {pm("Recall (micro)", "baseline")} / {pm("F1 (micro)", "baseline")}', f'{pm("Precision (micro)", "textcnn")} / {pm("Recall (micro)", "textcnn")} / {pm("F1 (micro)", "textcnn")}', f'{pm("Precision (micro)", "bilstm")} / {pm("Recall (micro)", "bilstm")} / {pm("F1 (micro)", "bilstm")}', f'{sv(sb, "Precision_micro")} / {sv(sb, "Recall_micro")} / {sv(sb, "F1_micro")}', f'{sv(sg, "Precision_micro")} / {sv(sg, "Recall_micro")} / {sv(sg, "F1_micro")}', f'{sv(stf, "Precision_micro")} / {sv(stf, "Recall_micro")} / {sv(stf, "F1_micro")}'],
        ["Precision / recall / F1, weighted", f'{pm("Precision (weighted)", "baseline")} / {pm("Recall (weighted)", "baseline")} / {pm("F1 (weighted)", "baseline")}', f'{pm("Precision (weighted)", "textcnn")} / {pm("Recall (weighted)", "textcnn")} / {pm("F1 (weighted)", "textcnn")}', f'{pm("Precision (weighted)", "bilstm")} / {pm("Recall (weighted)", "bilstm")} / {pm("F1 (weighted)", "bilstm")}', f'{sv(sb, "Precision_weighted")} / {sv(sb, "Recall_weighted")} / {sv(sb, "F1_weighted")}', f'{sv(sg, "Precision_weighted")} / {sv(sg, "Recall_weighted")} / {sv(sg, "F1_weighted")}', f'{sv(stf, "Precision_weighted")} / {sv(stf, "Recall_weighted")} / {sv(stf, "F1_weighted")}'],
        ["Macro-F1 95% CI", f'{pm("Macro-F1 95% CI low", "baseline")} to {pm("Macro-F1 95% CI high", "baseline")}', f'{pm("Macro-F1 95% CI low", "textcnn")} to {pm("Macro-F1 95% CI high", "textcnn")}', f'{pm("Macro-F1 95% CI low", "bilstm")} to {pm("Macro-F1 95% CI high", "bilstm")}', f'{sv(sb, "F1_macro_CI95_low")} to {sv(sb, "F1_macro_CI95_high")}', f'{sv(sg, "F1_macro_CI95_low")} to {sv(sg, "F1_macro_CI95_high")}', f'{sv(stf, "F1_macro_CI95_low")} to {sv(stf, "F1_macro_CI95_high")}'],
        ["Confusion TN / FP / FN / TP", f'{pw["Confusion TN"]["baseline"]} / {pw["Confusion FP"]["baseline"]} / {pw["Confusion FN"]["baseline"]} / {pw["Confusion TP"]["baseline"]}', f'{pw["Confusion TN"]["textcnn"]} / {pw["Confusion FP"]["textcnn"]} / {pw["Confusion FN"]["textcnn"]} / {pw["Confusion TP"]["textcnn"]}', f'{pw["Confusion TN"]["bilstm"]} / {pw["Confusion FP"]["bilstm"]} / {pw["Confusion FN"]["bilstm"]} / {pw["Confusion TP"]["bilstm"]}', f'{sb["CM_TN"]} / {sb["CM_FP"]} / {sb["CM_FN"]} / {sb["CM_TP"]}', f'{sg["CM_TN"]} / {sg["CM_FP"]} / {sg["CM_FN"]} / {sg["CM_TP"]}', f'{stf["CM_TN"]} / {stf["CM_FP"]} / {stf["CM_FN"]} / {stf["CM_TP"]}'],
        ["ROC-AUC", pm("ROC-AUC", "baseline"), pm("ROC-AUC", "textcnn"), pm("ROC-AUC", "bilstm"), sv(sb, "ROC_AUC"), sv(sg, "ROC_AUC"), sv(stf, "ROC_AUC")],
        ["PR-AUC", pm("PR-AUC", "baseline"), pm("PR-AUC", "textcnn"), pm("PR-AUC", "bilstm"), sv(sb, "PR_AUC"), sv(sg, "PR_AUC"), sv(stf, "PR_AUC")],
        ["MCC", pm("MCC", "baseline"), pm("MCC", "textcnn"), pm("MCC", "bilstm"), sv(sb, "MCC"), sv(sg, "MCC"), sv(stf, "MCC")],
        ["MCC 95% CI", f'{pm("MCC 95% CI low", "baseline")} to {pm("MCC 95% CI high", "baseline")}', f'{pm("MCC 95% CI low", "textcnn")} to {pm("MCC 95% CI high", "textcnn")}', f'{pm("MCC 95% CI low", "bilstm")} to {pm("MCC 95% CI high", "bilstm")}', f'{sv(sb, "MCC_CI95_low")} to {sv(sb, "MCC_CI95_high")}', f'{sv(sg, "MCC_CI95_low")} to {sv(sg, "MCC_CI95_high")}', f'{sv(stf, "MCC_CI95_low")} to {sv(stf, "MCC_CI95_high")}'],
        ["Brier score", pm("Brier Score", "baseline"), pm("Brier Score", "textcnn"), pm("Brier Score", "bilstm"), sv(sb, "Brier"), sv(sg, "Brier"), sv(stf, "Brier")],
        ["Expected calibration error", pm("ECE", "baseline"), pm("ECE", "textcnn"), pm("ECE", "bilstm"), sv(sb, "ECE"), sv(sg, "ECE"), sv(stf, "ECE")],
        ["McNemar vs own baseline: b / c / statistic / p", "", f'{fmt(pw["McNemar b (baseline right, model wrong)"]["textcnn"], 0)} / {fmt(pw["McNemar c (baseline wrong, model right)"]["textcnn"], 0)} / {pm("McNemar Statistic", "textcnn", 2)} / {pm("McNemar p-value", "textcnn", 3)}', f'{fmt(pw["McNemar b (baseline right, model wrong)"]["bilstm"], 0)} / {fmt(pw["McNemar c (baseline wrong, model right)"]["bilstm"], 0)} / {pm("McNemar Statistic", "bilstm", 2)} / {pm("McNemar p-value", "bilstm", 3)}', "", f'{sv(sg, "McNemar_b", 0)} / {sv(sg, "McNemar_c", 0)} / {sv(sg, "McNemar_stat", 2)} / {sv(sg, "McNemar_p", 3)}', f'{sv(stf, "McNemar_b", 0)} / {sv(stf, "McNemar_c", 0)} / {sv(stf, "McNemar_stat", 2)} / {sv(stf, "McNemar_p", 3)}'],
        ["Best epoch", pw["Best Epoch"]["baseline"], pw["Best Epoch"]["textcnn"], pw["Best Epoch"]["bilstm"], sb["Best_Epoch"], sg["Best_Epoch"], stf["Best_Epoch"]],
        ["Training time (sec)", pm("Training Time (sec)", "baseline", 1), pm("Training Time (sec)", "textcnn", 1), pm("Training Time (sec)", "bilstm", 1), sv(sb, "Train_Time_s", 1), sv(sg, "Train_Time_s", 1), sv(stf, "Train_Time_s", 1)],
        ["Training examples/sec", pm("Training Examples/sec", "baseline", 0), pm("Training Examples/sec", "textcnn", 0), pm("Training Examples/sec", "bilstm", 0), sv(sb, "Train_Examples_per_s", 0), sv(sg, "Train_Examples_per_s", 0), sv(stf, "Train_Examples_per_s", 0)],
        ["Inference examples/sec", pm("Inference Examples/sec", "baseline", 0), pm("Inference Examples/sec", "textcnn", 0), pm("Inference Examples/sec", "bilstm", 0), sv(sb, "Inference_Examples_per_s", 0), sv(sg, "Inference_Examples_per_s", 0), sv(stf, "Inference_Examples_per_s", 0)],
        ["Peak memory (MB)", pm("Peak Memory (MB)", "baseline", 1), pm("Peak Memory (MB)", "textcnn", 1), pm("Peak Memory (MB)", "bilstm", 1), sv(sb, "Peak_Memory_MB", 1), sv(sg, "Peak_Memory_MB", 1), sv(stf, "Peak_Memory_MB", 1)],
    ]
    r.table(header, rows, font=7, widths=[1.4, 1.35, 1.35, 1.35, 1.35, 1.35, 1.35])

    r.h("4.2 Per-slice macro-F1 / error rate", 2)
    slice_rows = []
    for label, key in [("Short reviews (P: under 50 words)", "length_bucket=short"), ("Medium reviews (P: 50 to 150 words)", "length_bucket=medium"), ("Long reviews (P: over 150 words)", "length_bucket=long"), ("Contains negation", "has_negation=1"), ("No negation", "has_negation=0"), ("Contains 'but'", "has_but=1"), ("No 'but'", "has_but=0"), ("Contains an exclamation mark", "has_exclamation=1"), ("No exclamation mark", "has_exclamation=0")]:
        slice_rows.append([label, pw[f"Support [{key}]"]["baseline"]] + [f'{fmt(pw[f"Macro-F1 [{key}]"][m], 3)} / {pct(pw[f"Error Rate [{key}]"][m])}' for m in ("baseline", "textcnn", "bilstm")])
    r.table(["Poushali's slices", "N", "baseline", "textcnn", "bilstm"], slice_rows, font=8)
    r.table(["Sadaf's slices", "N", "Baseline", "BiGRU + attention", "Transformer encoder"], [
        ["Short (up to 50 words)", "9,362", "0.900 / 9.6%", "0.932 / 6.5%", "0.918 / 7.9%"],
        ["Medium (51 to 149 words)", "16,853", "0.906 / 9.4%", "0.937 / 6.3%", "0.923 / 7.7%"],
        ["Long (150 words and more)", "11,785", "0.897 / 9.9%", "0.927 / 7.1%", "0.917 / 8.0%"],
        ["Contains negation", "28,514", "0.896 / 10.1%", "0.930 / 6.9%", "0.914 / 8.3%"],
        ["Truncated (over 256 tokens)", "782", "0.856 / 11.8%", "0.881 / 10.2%", "0.900 / 8.2%"],
    ], font=8)
    r.portrait()
    r.picture(P2 / "outputs" / "full" / "bilstm" / "confusion_matrix.png", 3.2, "Poushali: BiLSTM confusion matrix on the 38,000 test reviews.")
    r.picture(S2 / "outputs" / "confusion_matrices.png", 6.0, "Sadaf: confusion matrices of her three models.")

    r.h("4.3 Joint analysis", 2)
    r.p(
        "Strengths. Six models trained without pretrained components on the same official test set, evaluated with the same metric list, bootstrap intervals, and paired McNemar tests, give a clear and defensible ranking. "
        "Sadaf's BiGRU with attention pooling is the best model in the team at 93.4% accuracy (macro-F1 0.934, MCC 0.869, ROC-AUC 0.984), followed by Poushali's BiLSTM at 93.1% (MCC 0.861, ROC-AUC 0.981). Both recurrent models beat their own baselines significantly (McNemar p about 1e-87 and 1e-9). "
        "Calibration is good across the board: every model has ECE below 0.025, and the two simplest baselines are the best calibrated (0.004 and 0.006)."
    )
    r.p(
        "Weaknesses. Reading word order matters, but how much depends on the baseline it is compared with: Poushali's mean-pooled MLP baseline already reaches 92.3%, so her BiLSTM adds 0.7 points, while Sadaf's max-pooled linear baseline sits at 90.4%, so her BiGRU adds 3.0 points. Poushali's TextCNN is not a significant improvement over her baseline (p 0.12). Sadaf's from-scratch Transformer encoder (92.2%) trails both recurrent models, as small Transformers usually do on 100K examples, but it is the best model on truncated reviews. "
        "Both error reviews find the same failure families: mixed reviews whose verdict comes late or after a 'but', temporal reversals ('used to be terrible, now much better'), sarcasm, and a visible share of label noise. Both slice analyses agree that reviews containing negation are the hardest and that the recurrent models close most of that gap."
    )
    r.p(
        "Limitations. Each member used 100K of the 560K training reviews to fit the lab schedule; stemming (Poushali) and lemmatization (Sadaf) both discard morphology that carries sentiment; stopword removal drops intensifiers; 256-token truncation cuts about 2% of reviews, exactly the ones where the verdict is most often at the end. The two members' slices are defined slightly differently (word-count thresholds and the truncated slice), so slice numbers are compared within a member, not across."
    )
    r.p(
        "What the team would try next. Train on the full 560K reviews; add attention pooling to the BiLSTM, since attention pooling is the one design difference that correlates with Sadaf's higher score; keep 'but', 'since', 'now', and intensifiers in the vocabulary; give the recurrent models an explicit view of the closing sentences, the fix proposed and made testable in Poushali's error review; and raise the token limit or process long reviews hierarchically, the fix proposed in Sadaf's."
    )

    r.h("4.4 Error reviews, 20 cases each", 2)
    r.p("Poushali Purkayastha, BiLSTM (checkpoints/full/bilstm/best.pt). Full text of every review is in task2_sentiment/Poushali_Purkayastha/outputs/full/bilstm/error_review_candidates.md; the complete write-up is in her failure_analysis.md. Label 1 is positive.")
    r.table(["#", "Group", "Test index", "Label / pred / p(pos)", "Snippet", "Error type"], [
        ["1", "confident FP", "25815", "0 / 1 / 1.000", "The TrimTini I had was really delicious... it's not likely I will return any time soon. But it was a nice event, and I enjoyed myself.", "mixed sentiment"],
        ["2", "confident FP", "17756", "0 / 1 / 1.000", "I absolutely love how it is decorated... The gel manicure itself wasn't done well at all... I can find better service and product for cheaper", "mixed sentiment"],
        ["3", "confident FP", "9609", "0 / 1 / 1.000", "Sure it has... a terrific buffet... But the hotel and casino is not good for one reason only -- SMOKERS... Highly recommended.", "sarcasm or irony"],
        ["4", "confident FP", "8945", "0 / 1 / 1.000", "Bianco's is the best marketing ploy in the city... you have to arrive an hour before the restaurant even opens to wait in line...", "sarcasm or irony"],
        ["5", "confident FP", "30899", "0 / 1 / 1.000", "the room was exactly what we needed... we did not get any sleep... Besides that the room was nice", "mixed sentiment"],
        ["6", "confident FN", "22807", "1 / 0 / 0.000", "EDIT: They really did change the service up since I last posted this. Horrible service.", "likely label noise"],
        ["7", "confident FN", "30793", "1 / 0 / 0.000", "so much better since they changed owners... it was terrible... Now its much better.", "temporal reversal"],
        ["8", "confident FN", "11401", "1 / 0 / 0.000", "I was a little disappointed... wrong food and then cold food... head there for drinks but not for food.", "mixed sentiment"],
        ["9", "confident FN", "19296", "1 / 0 / 0.000", "the sweet and sour beef looked like fried spam... soooo ewwwwww!!... This place is good for groups, but service is slow.", "mixed sentiment"],
        ["10", "confident FN", "30366", "1 / 0 / 0.000", "no acknowledgement of a birthday... absolutely no offer of a gift card... That was not good business.", "likely label noise"],
        ["11", "near threshold", "32036", "0 / 1 / 0.500", "This place is not only not open during business hours but looks like it's shuttered it's doors for good...", "negation or contrast scope"],
        ["12", "near threshold", "37885", "1 / 0 / 0.499", "this one by far has no good items. Couldn't find anything good here.", "likely label noise"],
        ["13", "near threshold", "20958", "0 / 1 / 0.501", "we realized there was a fire... we never got anything in the mail... we love it. But till the time our issue is resolved we will not go", "narrative with weak sentiment cues"],
        ["14", "near threshold", "10954", "0 / 1 / 0.502", "I would like to give this place 1 star only but... Love that it's 24 hours... He should be FIRED!... I'll give them one more star :)", "sarcasm or irony"],
        ["15", "near threshold", "11883", "1 / 0 / 0.498", "This museum is a BLAST! It's THE BOMB!... the first one to use those lame puns?", "figurative language or slang"],
        ["16", "slice: no exclamation", "25701", "0 / 1 / 1.000", "the gorditas were filling and good.....but not great... there are much better options... The service was very friendly", "mixed sentiment"],
        ["17", "slice: no exclamation", "15988", "0 / 1 / 1.000", "I find their meats to be generally good... an enjoyable adventure. Prices are downright cheap.", "likely label noise"],
        ["18", "slice: no exclamation", "28036", "0 / 1 / 1.000", "you provide one pickle on a plate? Amazing... Better than freshly made sliders without pickles. Ahhhhhhhhhhhhh.", "sarcasm or irony"],
        ["19", "slice: no exclamation", "4488", "0 / 1 / 1.000", "It's good. The rolls are better than the sashimi... Little Tokyo and Kiku are better... 4.5 of 10 relative to sushi quality alone", "numeric cue lost in preprocessing"],
        ["20", "slice: no exclamation", "4806", "1 / 0 / 0.000", "I usually associate anything with the government or DMV as a huge pain..., but this was extremely fast and easy... no doubt that I will go back", "negation or contrast scope"],
    ], font=7, widths=[0.3, 0.9, 0.6, 0.9, 3.2, 1.2])
    r.p("Summary: mixed sentiment 6, sarcasm 4, likely label noise 4, negation or contrast scope 2, temporal reversal 1, figurative language 1, numeric cue lost 1, weak-cue narrative 1. Proposed testable fix: concatenate the BiLSTM's final states with a max-pool over the last 32 valid tokens and keep 'but', 'since', 'now' in the vocabulary; success is a macro-F1 gain of at least 0.5 points on the 'contains but' slice with no loss on the 'no but' slice and a significant paired McNemar test against the current BiLSTM.")

    r.p("Sadaf Fatima Syeda, BiGRU with attention (checkpoints/BiGRU_Attention.pt). From task2_sentiment/Sadaf_Fatima_Syeda/failure_analysis.md; the source file is outputs/error_review_BiGRU_Attention.csv.")
    r.table(["#", "Group", "Snippet", "p(pos)", "Error type", "Proposed fix"], [
        ["1", "confident FP", "Wow love the place and everything is very clean and new! Great place to come and relax worth a try!", "1.000", "positive lexical cues or possible label noise", "more balanced examples; inspect similar training reviews"],
        ["2", "confident FP", "The service was awesome; space impressive. But... that NYC gem looses something and just isn't the same.", "0.9998", "contrast and negation", "contrastive examples with 'but', 'however', 'isn't'"],
        ["3", "confident FP", "20 years for me and sad to see them go... consistently great fresh fish... END OF AN ERA.", "0.9998", "mixed sentiment, nostalgic wording", "more examples where praise and disappointment occur together"],
        ["4", "confident FP", "The food is a tad better than okay... The manager is awesome... nerve wrecking.", "0.9996", "mixed sentiment, negative final judgment", "pooling that weights the overall conclusion"],
        ["5", "confident FP", "I absolutely love how it is decorated... The gel manicure itself wasn't done well at all.", "0.9989", "aspect-level sentiment", "aspect-aware examples"],
        ["6", "confident FN", "The beer is Delicious and so is the food - However I unfortunately, can not say the same about the HELP. The service was terrible", "0.0001", "negation and contrast", "more examples with 'however', 'not', 'terrible'"],
        ["7", "confident FN", "This review is for the pharmacy only... Price difference is unbelievable!! ...No wonder premiums are high", "0.0001", "negative-sounding words in a positive review", "positive reviews with negative-sounding or domain-specific language"],
        ["8", "confident FN", "Pizza is great. However, the service is not... UPDATE: service has improved drastically.", "0.0002", "temporal update", "examples with updates and temporal changes"],
        ["9", "confident FN", "The children were running in and out of the bar... the bartender was very cordial", "0.0002", "long mixed-sentiment review", "longer context or hierarchical modelling"],
        ["10", "confident FN", "This place is so much better since they changed owners... it was terrible... Now its much better.", "0.0006", "temporal contrast", "examples that distinguish past and present sentiment"],
        ["11", "near threshold", "This place was horrible... Insipid, over-battered... On the bright side, service was very fast.", "0.5006", "mixed sentiment", "mixed-review examples, aspect-level attention"],
        ["12", "near threshold", "It was total bootleg up in this location... So glad there are other locations of Ross stores near by offering the same great deals.", "0.5006", "sarcasm and indirect sentiment", "sarcastic and humorous reviews in training"],
        ["13", "near threshold", "The first time... was a 3-star experience... Last week's experience was a real let down... the setting is very pretty", "0.5007", "conflicting evidence across visits", "temporal and aspect-aware examples"],
        ["14", "near threshold", "Walmart makes me cringe... But... The pharmacy staff have always been pleasant and efficient.", "0.5008", "mixed sentiment across parts of the business", "separate sentiment by aspect"],
        ["15", "near threshold", "Fast no wait seating on a Saturday night ...but be prepared to pay for substandard food!", "0.5012", "contrast between quick service and poor food", "contrastive examples"],
        ["16", "slice: truncated", "Hotel review: long list of problems, ends 'The entire property is nice, which I would stay there again'", "0.0033", "positive ending cut off by truncation", "increase the context length or hierarchical processing"],
        ["17", "slice: truncated", "Long humorous casino story with very little about the venue itself", "0.0063", "weak direct sentiment cues", "longer humorous reviews, document-level context"],
        ["18", "slice: truncated", "Waste facility visit: 'the place stinks... bins were full', ends 'Can't beat that!'", "0.0072", "negative details, positive or humorous ending", "preserve the full review"],
        ["19", "slice: truncated", "Show review: 'worst show you will have ever seen' jokes, ends 'top 2-3 shows in town... Totally would recommend it'", "0.0076", "sarcasm and joke-negative wording", "sarcastic reviews, longer context window"],
        ["20", "slice: truncated", "Restaurant review: first visit great, then 'concept change' disappointed, 'it will probably be the last'", "0.9899", "positive descriptions outweigh the negative judgment", "longer sequence length, temporal examples"],
    ], font=7, widths=[0.3, 0.9, 3.0, 0.5, 1.3, 1.3])
    r.p("Summary: mixed sentiment, negation or contrast, temporal updates, sarcasm, and loss of context in long reviews. Proposed fix: a longer context window or hierarchical modelling, tested by retraining the BiGRU with a longer window and comparing overall accuracy and the truncated-review slice against the current model.")

    # ---------------- Task 3 ----------------
    p3 = kv_csv(P3 / "full_metrics_report.csv")
    s3 = group_csv(S3 / "full_metrics_report.csv")
    ko = json.loads((P3 / "outputs" / "full" / "kaggle" / "official_scores.json").read_text(encoding="utf-8"))
    kr = json.loads((P3 / "outputs" / "full" / "kaggle" / "kaggle_results.json").read_text(encoding="utf-8"))
    skr = json.loads((S3 / "outputs" / "kaggle_results.json").read_text(encoding="utf-8"))
    r.h("5. Task 3: CycleGAN Monet and photo style transfer", 1)
    r.h("5.1 Comparison table", 2)
    r.p("Directions are named by content: photo to Monet is the Kaggle direction. In Sadaf's files domain A is Monet, so her 'B2A' values are photo to Monet; in Poushali's files domain A is photo.")
    rows = [
        ["Generators", "ResNet: 7x7 conv 64, two stride-2 convs to 256, 6 residual blocks, two transposed convs, 7x7 conv, tanh; reflection padding, instance norm; 7,837,699 parameters each", "U-Net: 8 downsampling levels, 48 base channels, max multiplier 8; 30,604,035 parameters each"],
        ["Discriminators", "70x70 PatchGAN, 4 conv layers 64 to 512, instance norm, LeakyReLU 0.2; 2,764,737 parameters each", "PatchGAN with 2 layers, ndf 64; 662,593 parameters each"],
        ["Adversarial loss, cycle, identity", "least-squares; cycle 10; identity 5", "sigmoid cross-entropy; cycle 8; no identity"],
        ["Image size, batch, steps", "128 px training, 256 px inference; batch 4; 7,200 steps (24 epochs x 300)", "256 px (load 286 with random crop); batch 1; 20,000 steps"],
        ["Optimizer, schedule", "Adam 2e-4, beta1 0.5, constant 12 epochs then linear decay", "Adam 2e-4, betas 0.5/0.999"],
        ["Holdout for paired metrics", "300 photos, 30 Monets (seed 20)", "300 images (seed 266)"],
        ["Hardware, mixed precision", "RTX 5090, bfloat16", "Tesla T4, AMP"],
        ["Training time, images/sec", f'{fmt(p3["Training time (sec)"], 1)} s, {fmt(p3["Training images/sec"], 1)}', f'{fmt(float(s3["training time (min)"]) * 60, 1)} s, {fmt(s3["training images/sec"], 1)}'],
        ["Peak training memory (MB)", fmt(p3["Peak memory (MB)"], 1), fmt(s3["peak memory training (MB)"], 1)],
        ["Parameters total", fmt(p3["Parameter count total"]), fmt(s3["parameters total"])],
        ["FID, photo to Monet", fmt(p3["FID photo->monet"], 2), fmt(s3["FID Photo->Monet"], 2)],
        ["FID, Monet to photo", fmt(p3["FID monet->photo"], 2), fmt(s3["FID Monet->Photo"], 2)],
        ["KID mean (std), photo to Monet", f'{fmt(p3["KID mean photo->monet"])} ({fmt(p3["KID std photo->monet"])})', f'{fmt(s3["KID Photo->Monet"])} ({fmt(s3["KID std Photo->Monet"])})'],
        ["KID mean (std), Monet to photo", f'{fmt(p3["KID mean monet->photo"])} ({fmt(p3["KID std monet->photo"])})', f'{fmt(s3["KID Monet->Photo"])} ({fmt(s3["KID std Monet->Photo"])})'],
        ["Precision / recall, photo to Monet", f'{fmt(p3["Precision photo->monet"], 3)} / {fmt(p3["Recall photo->monet"], 3)}', f'{fmt(s3["precision Photo->Monet"], 3)} / {fmt(s3["recall Photo->Monet"], 3)}'],
        ["Density / coverage, photo to Monet", f'{fmt(p3["Density photo->monet"], 3)} / {fmt(p3["Coverage photo->monet"], 3)}', f'{fmt(s3["density Photo->Monet"], 3)} / {fmt(s3["coverage Photo->Monet"], 3)}'],
        ["Precision / recall, Monet to photo", f'{fmt(p3["Precision monet->photo"], 3)} / {fmt(p3["Recall monet->photo"], 3)}', f'{fmt(s3["precision Monet->Photo"], 3)} / {fmt(s3["recall Monet->Photo"], 3)}'],
        ["Density / coverage, Monet to photo", f'{fmt(p3["Density monet->photo"], 3)} / {fmt(p3["Coverage monet->photo"], 3)}', f'{fmt(s3["density Monet->Photo"], 3)} / {fmt(s3["coverage Monet->Photo"], 3)}'],
        ["Cycle-reconstruction L1, photo side / Monet side", f'{fmt(p3["Cycle-reconstruction L1 photo->monet->photo"])} / {fmt(p3["Cycle-reconstruction L1 monet->photo->monet"])}', f'{fmt(s3["cycle-reconstruction L1 Photo->Monet (0-1 scale)"])} / {fmt(s3["cycle-reconstruction L1 Monet->Photo (0-1 scale)"])}'],
        ["Identity L1, photos / Monets", f'{fmt(p3["Identity L1 G_BA(photo) vs photo"])} / {fmt(p3["Identity L1 G_AB(monet) vs monet"])} (trained)', f'{fmt(s3["identity L1 Photo->Monet (0-1 scale, not trained)"])} / {fmt(s3["identity L1 Monet->Photo (0-1 scale, not trained)"])} (not trained)'],
        ["LPIPS input vs translation, photo to Monet / Monet to photo", f'{fmt(p3["LPIPS input vs translation photo->monet"], 3)} / {fmt(p3["LPIPS input vs translation monet->photo"], 3)}', f'{fmt(s3["LPIPS input vs translation Photo->Monet"], 3)} / {fmt(s3["LPIPS input vs translation Monet->Photo"], 3)}'],
        ["LPIPS input vs reconstruction, photo side / Monet side", f'{fmt(p3["LPIPS input vs reconstruction photo->monet->photo"], 3)} / {fmt(p3["LPIPS input vs reconstruction monet->photo->monet"], 3)}', f'{fmt(s3["LPIPS input vs reconstruction Photo->Monet"], 3)} / {fmt(s3["LPIPS input vs reconstruction Monet->Photo"], 3)}'],
        ["Content cosine similarity, photo to Monet / Monet to photo", f'{fmt(p3["Content cosine similarity photo->monet"], 3)} / {fmt(p3["Content cosine similarity monet->photo"], 3)}', f'{fmt(s3["content-preservation cosine Photo->Monet"], 3)} / {fmt(s3["content-preservation cosine Monet->Photo"], 3)}'],
        ["Final generator / discriminator losses", f'G {fmt(p3["Final epoch mean loss_g"], 3)}, D_A {fmt(p3["Final epoch mean loss_d_a"], 3)}, D_B {fmt(p3["Final epoch mean loss_d_b"], 3)}, cycle {fmt(p3["Final epoch mean loss_cycle"], 3)}, identity {fmt(p3["Final epoch mean loss_identity"], 3)}', f'G {fmt(s3["G_total loss"], 3)}, D_A {fmt(s3["D_A loss"], 3)}, D_B {fmt(s3["D_B loss"], 3)}, cycle {fmt(s3["cycle loss"], 3)}, identity {fmt(s3["identity loss"], 3)}'],
        ["Gradient norms, NaN count", f'G mean {fmt(p3["Gradient norm G mean"], 1)} max {fmt(p3["Gradient norm G max"], 1)}; D mean {fmt(p3["Gradient norm D mean"], 1)} max {fmt(p3["Gradient norm D max"], 1)}; NaN {p3["NaN or Inf losses"]}', f'G max {s3["gnorm_G_max"]}, D max {s3["gnorm_D_max"]} (fp16 overflow readings); non-finite steps {s3["non-finite (NaN/Inf) steps"]}'],
        ["Competition FID / MiFID / score (submission.csv)", f'{fmt(ko["FID"], 2)} / {fmt(ko["MiFID"], 4)} / {fmt(ko["score"], 3)}, all 7,038 generated images against the competition real_stats.npz', f'{fmt(s3["submission FID (mean of both directions)"], 2)} / {fmt(s3["submission MiFID (mean of both directions)"], 4)} / {fmt(s3["leaderboard score estimate (FID + MiFID) / 2"], 3)}, mean of both directions on 300 images'],
        ["Kaggle public score, date", f'{kr["public_score"]}, 2026-09-28 (rank {kr["rank"]} that day)', f'{skr["public_score"]}, {skr["submitted_on"]}'],
        ["Human audit", "pending: blinded 30-photo set built with src/audit.py, both raters to score style, content, artifacts", "pending: blinded set in outputs/audit with sheets for both raters"],
        ["Evidence", "task3_gan/Poushali_Purkayastha/outputs/full: loss_curves.png, lr_schedule.png, samples/, pred_A2B_preview, pred_B2A_preview, kaggle/; checkpoints/full/generators.pt; raw log and manifest Poushali_Purkayastha_task3_gan_full_20260928_180436", "task3_gan/Sadaf_Fatima_Syeda/outputs: loss_curves.png, final_grid.png, samples/, pred_A2B, pred_B2A; checkpoints/G_AB_monet2photo_fp16.pt, G_BA_photo2monet_fp16.pt; raw log and manifest task3_sadaf_20260927_181844"],
    ]
    r.table(["", "Poushali Purkayastha", "Sadaf Fatima Syeda"], rows, font=8, widths=[1.7, 2.55, 2.55])
    r.picture(P3 / "outputs" / "full" / "loss_curves.png", 6.4, "Poushali: generator and discriminator losses, loss components, and gradient norms over 7,200 steps.")
    r.picture(P3 / "outputs" / "full" / "samples" / "epoch_024.jpg", 3.4, "Poushali, epoch 24 sample grid. Rows: real photo, fake Monet, reconstructed photo, real Monet, fake photo, reconstructed Monet.")
    r.picture(S3 / "outputs" / "loss_curves.png", 6.0, "Sadaf: training losses over 20,000 steps.")
    r.picture(S3 / "outputs" / "final_grid.png", 5.0, "Sadaf: final sample grid.")

    r.h("5.2 Kaggle leaderboard", 2)
    r.p(
        f"Both members submitted under PairProgramming_Team_16 in the required format, a submission.csv with ID, FID, and MiFID. Poushali's submission scored {kr['public_score']} and Sadaf's {skr['public_score']}; lower is better, and the score is the average of FID and MiFID as defined on the competition's Overview page. "
        "Poushali's numbers were computed by src/kaggle_eval.py on all 7,038 generated images against the competition's real_stats.npz, whose stored features match ours with cosine 0.9999, so they follow the course definitions exactly. "
        "On 2026-09-28 the team's best submission ranked first on the public leaderboard. The private score and the final rank will be recorded in each member's kaggle_results.json and metrics report when the competition closes. Every submitted image is the unedited output of the member's own generator; the full image sets are archived on Google Drive."
    )

    r.h("5.3 Joint analysis", 2)
    r.p(
        "Strengths. Both CycleGANs learned the task: cycle reconstructions are close to their inputs (cycle L1 of 0.044 to 0.081 on a 0 to 1 scale for both members), content is preserved through translation (Inception cosine similarity 0.66 to 0.79), and neither run diverged. "
        "Poushali's ResNet generator with least-squares loss and identity loss reached the better competition score (48.86 against 66.90) and the better FID in the Kaggle direction (97.3 against 123.3), with a KID of 0.021 against 0.029, in a 9-minute run on the RTX 5090. "
        "Sadaf's U-Net reached the lower cycle L1 (0.044 against 0.064) and, in the Monet-to-photo direction, better precision and coverage, and it was trained at the native 256 px resolution."
    )
    r.p(
        "Weaknesses. Poushali's model shows a checkerboard texture in its 256 px outputs, a consequence of transposed-convolution upsampling in a generator trained at 128 px, and it translates Monet to photo weakly (LPIPS 0.24, coverage 0.15), which the identity loss and the 270-image Monet domain explain. "
        "Sadaf's model changes its inputs more (LPIPS 0.54 to 0.56 against 0.24 to 0.36) at the cost of realism as measured by FID, and without an identity term its untrained identity L1 is high (0.33 to 0.43), meaning the generators recolour images that already belong to the target domain. Her discriminator losses stayed near 0.5 throughout, the balanced regime of BCE training, while Poushali's Monet discriminator pulled ahead during learning-rate decay."
    )
    r.p(
        "Limitations. FID against only 300 real Monets is noisy and upward-biased for every team, so absolute values should not be compared with published numbers; KID and the manifold metrics are reported for that reason. The two members' competition scores are not computed on identical image sets (all 7,038 photos against a 300-image holdout, and one direction against the mean of both), so the leaderboard, which scores the same way for everyone, is the fair comparison. The human audit had not been scored when this report was written, so image quality is judged here only through Inception features and the sample grids."
    )
    r.p(
        "What the team would try next. Combine the two designs: the ResNet generator with least-squares loss at 256 px training, with identity loss only on the photo-to-Monet generator; replace transposed convolutions with upsample-and-convolution to remove the checkerboard; halve the discriminator learning rate and train for 60 or more epochs, which the lab benchmark shows is affordable; and finish the blinded audit so a human judgment sits beside the automated metrics."
    )

    r.h("5.4 Failure and artifact analyses", 2)
    r.p("Poushali Purkayastha, from task3_gan/Poushali_Purkayastha/failure_analysis.md.")
    r.table(["Case", "Where to look", "Failure type", "Observation"], [
        ["1", "outputs/full/pred_A2B_preview/000ded5c41.jpg (beach at sunset)", "checkerboard or tiling artifacts", "A regular fine grid of light and dark pixels over the sky and wet sand, with the period of the two transposed-convolution layers, stronger at 256 px inference than in the 128 px training grids; the 70x70 PatchGAN cannot penalize texture finer than its receptive field."],
        ["2", "outputs/full/samples/epoch_024.jpg row 2 column 1 (church against sky); pred_A2B_preview/00068bc07f.jpg", "colour speckle, hallucinated texture", "Flat sky regions are filled with purple, blue, and orange blotches absent from both source and Monets; the discriminator rewards colour variation everywhere and the cycle loss does not object because G_BA removes the speckle on the way back."],
        ["3", "outputs/full/samples/epoch_024.jpg rows 4 and 5; pred_B2A_preview/000c1e3bff.jpg", "unchanged output, near identity", "Monet-to-photo outputs keep brush strokes and painted skies; LPIPS 0.24 and coverage 0.15 confirm G_BA changes little, because the Monet domain has 270 training images against 6,738 photos and identity loss holds G_BA close to the identity."],
    ], font=8, widths=[0.4, 2.0, 1.3, 3.2])
    r.p("Training stability: no NaN or Inf loss in 7,200 steps, mixed precision on throughout, nine gradient-norm spikes above 100 in the first three epochs clipped at 10, and a Monet discriminator that gained the upper hand during decay (D_B loss 0.46 to 0.075 while G_AB adversarial loss rose 0.47 to 0.66).")
    r.picture(P3 / "outputs" / "full" / "pred_A2B_preview" / "000ded5c41.jpg", 2.4, "Poushali, case 1: checkerboard texture in a 256 px photo-to-Monet translation.")
    r.picture(P3 / "outputs" / "full" / "pred_B2A_preview" / "000c1e3bff.jpg", 2.4, "Poushali, case 3: a Monet-to-photo translation that remains a painting.")
    r.p("Sadaf Fatima Syeda: her Task 3 failure analysis is to be added to task3_gan/Sadaf_Fatima_Syeda/failure_analysis.md and to this section before the final push. Her sample grid (outputs/final_grid.png) and loss curves are included above, and her metrics record no non-finite steps and discriminator losses of about 0.5 in the last 20% of training.")

    r.h("5.5 Human audit", 2)
    r.p(
        "Protocol: 30 fixed holdout photos are translated by each member's generator, the outputs are shuffled and renamed so raters cannot tell which model produced them, and both members score every output from 1 to 5 on style (how Monet-like), content (how well the scene is kept), and artifacts (5 means none). Agreement is reported as Cohen's kappa, linear-weighted kappa, and percent agreement per criterion. "
        "Both members have built their blinded sets (task3_gan/Poushali_Purkayastha/src/audit.py and task3_gan/Sadaf_Fatima_Syeda/outputs/audit). The ratings and agreement values will be added to each member's metrics report and to this section when both sheets are filled."
    )

    r.h("6. References", 1)
    r.bullets([
        "Vaswani, A., Shazeer, N., Parmar, N., Uszkoreit, J., Jones, L., Gomez, A. N., Kaiser, L., Polosukhin, I. (2017). Attention Is All You Need. NeurIPS. https://arxiv.org/abs/1706.03762",
        "Eldan, R., Li, Y. (2023). TinyStories: How Small Can Language Models Be and Still Speak Coherent English? https://arxiv.org/abs/2305.07759",
        "Zhu, J.-Y., Park, T., Isola, P., Efros, A. A. (2017). Unpaired Image-to-Image Translation using Cycle-Consistent Adversarial Networks. ICCV. https://arxiv.org/abs/1703.10593",
        "Kim, Y. (2014). Convolutional Neural Networks for Sentence Classification. EMNLP.",
        "Mao, X., Li, Q., Xie, H., Lau, R. Y. K., Wang, Z., Smolley, S. P. (2017). Least Squares Generative Adversarial Networks. ICCV.",
        "Isola, P., Zhu, J.-Y., Zhou, T., Efros, A. A. (2017). Image-to-Image Translation with Conditional Adversarial Networks. CVPR.",
        "Heusel, M., Ramsauer, H., Unterthiner, T., Nessler, B., Hochreiter, S. (2017). GANs Trained by a Two Time-Scale Update Rule Converge to a Local Nash Equilibrium. NeurIPS.",
        "Binkowski, M., Sutherland, D. J., Arbel, M., Gretton, A. (2018). Demystifying MMD GANs. ICLR.",
        "Kynkaanniemi, T., Karras, T., Laine, S., Lehtinen, J., Aila, T. (2019). Improved Precision and Recall Metric for Assessing Generative Models. NeurIPS.",
        "Naeem, M. F., Oh, S. J., Uh, Y., Choi, Y., Yoo, J. (2020). Reliable Fidelity and Diversity Metrics for Generative Models. ICML.",
        "Zhang, R., Isola, P., Efros, A. A., Shechtman, E., Wang, O. (2018). The Unreasonable Effectiveness of Deep Features as a Perceptual Metric. CVPR.",
        "Zhang, X., Zhao, J., LeCun, Y. (2015). Character-level Convolutional Networks for Text Classification. NeurIPS. (source of the Yelp polarity dataset)",
    ])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    d.save(OUT)
    print("written", OUT)


if __name__ == "__main__":
    build()
